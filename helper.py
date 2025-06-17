import openpyxl
import logging
import os
from aiogram import Bot, Dispatcher, types
from aiogram.contrib.fsm_storage.memory import MemoryStorage
from aiogram.dispatcher import FSMContext
from aiogram.dispatcher.filters.state import State, StatesGroup
from datetime import datetime

# Настройка бота
bot = Bot(token="YOUR TOKEN")
storage = MemoryStorage()
dp = Dispatcher(bot, storage=storage)

# Настройка логирования
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# ID админов
ADMINS = [1585398371, 0]

# Клавиатуры
start_keyboard = types.ReplyKeyboardMarkup(resize_keyboard=True)
start_keyboard.add(types.KeyboardButton('/start'))

# Состояния
class UserQuestions(StatesGroup):
    waiting_for_organization = State()
    waiting_for_surname = State()
    waiting_for_name = State()
    waiting_for_question = State()

class AdminActions(StatesGroup):
    select_question = State()
    enter_answer = State()
    confirm_answer = State()

# Функции работы с Excel и пользовательскими данными
def init_excel():
    try:
        wb = openpyxl.load_workbook("questions.xlsx")
        return wb
    except FileNotFoundError:
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.append([
            "ID", "Организация", "Фамилия", "Имя", 
            "Время вопроса", "Вопрос", "Статус", 
            "Ответ", "ID админа", "Время ответа"
        ])
        wb.save("questions.xlsx")
        return wb

def init_user_data():
    if not os.path.exists("user_data.xlsx"):
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.append(["ID", "Организация", "Фамилия", "Имя"])
        wb.save("user_data.xlsx")

def get_user_data(user_id):
    try:
        wb = openpyxl.load_workbook("user_data.xlsx")
        sheet = wb.active
        
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row and row[0] == user_id:
                return {
                    'org': row[1],
                    'surname': row[2],
                    'name': row[3]
                }
        return None
    except Exception as e:
        logger.error(f"Error getting user data: {e}")
        return None

def save_user_data(user_id, organization, surname, name):
    try:
        wb = openpyxl.load_workbook("user_data.xlsx")
        sheet = wb.active
        
        # Проверяем, есть ли уже пользователь
        for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
            if row[0].value == user_id:
                # Обновляем данные
                row[1].value = organization
                row[2].value = surname
                row[3].value = name
                wb.save("user_data.xlsx")
                return
        
        # Если пользователя нет - добавляем
        sheet.append([user_id, organization, surname, name])
        wb.save("user_data.xlsx")
    except Exception as e:
        logger.error(f"Error saving user data: {e}")

def add_question(user_id, organization, surname, name, question):
    wb = init_excel()
    sheet = wb.active
    sheet.append([
        user_id, organization, surname, name,
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        question, "ожидает", None, None, None
    ])
    wb.save("questions.xlsx")

def get_pending_questions():
    try:
        wb = openpyxl.load_workbook("questions.xlsx")
        sheet = wb.active
        questions = []
        
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row and len(row) >= 7 and row[6] == "ожидает":
                questions.append({
                    'id': row[0],
                    'org': row[1],
                    'surname': row[2],
                    'name': row[3],
                    'time': row[4],
                    'question': row[5],
                    'user_info': f"{row[1]} ({row[3]} {row[2]})"
                })
        return questions
    except Exception as e:
        logger.error(f"Error getting questions: {e}")
        return []

def update_question(user_id, answer, admin_id):
    try:
        wb = openpyxl.load_workbook("questions.xlsx")
        sheet = wb.active
        
        for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
            if row[0].value == user_id and row[6].value == "ожидает":
                row[6].value = "Отвечено"
                row[7].value = answer
                row[8].value = admin_id
                row[9].value = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                wb.save("questions.xlsx")
                return True
        return False
    except Exception as e:
        logger.error(f"Error updating question: {e}")
        return False

# Инициализация файлов при запуске
init_excel()
init_user_data()

# Обработчики для пользователей
@dp.message_handler(commands=['start'])
async def start(message: types.Message):
    if message.from_user.id in ADMINS:
        await message.answer("Вы вошли как администратор. Используйте /answer для ответа на вопросы.")
    else:
        user_data = get_user_data(message.from_user.id)
        if user_data:
            await message.answer("Введите ваш вопрос:")
            await UserQuestions.waiting_for_question.set()
        else:
            await message.answer("Введите название вашей организации:")
            await UserQuestions.waiting_for_organization.set()

@dp.message_handler(state=UserQuestions.waiting_for_organization)
async def process_organization(message: types.Message, state: FSMContext):
    await state.update_data(organization=message.text)
    await message.answer("Введите вашу фамилию:")
    await UserQuestions.waiting_for_surname.set()

@dp.message_handler(state=UserQuestions.waiting_for_surname)
async def process_surname(message: types.Message, state: FSMContext):
    await state.update_data(surname=message.text)
    await message.answer("Введите ваше имя:")
    await UserQuestions.waiting_for_name.set()

@dp.message_handler(state=UserQuestions.waiting_for_name)
async def process_name(message: types.Message, state: FSMContext):
    data = await state.get_data()
    save_user_data(
        message.from_user.id,
        data['organization'],
        data['surname'],
        message.text
    )
    await message.answer("Введите ваш вопрос:")
    await UserQuestions.waiting_for_question.set()

@dp.message_handler(state=UserQuestions.waiting_for_question)
async def process_question(message: types.Message, state: FSMContext):
    user_data = get_user_data(message.from_user.id)
    if user_data:
        add_question(
            message.from_user.id,
            user_data['org'],
            user_data['surname'],
            user_data['name'],
            message.text
        )
    else:
        data = await state.get_data()
        add_question(
            message.from_user.id,
            data['organization'],
            data['surname'],
            data['name'],
            message.text
        )
        save_user_data(
            message.from_user.id,
            data['organization'],
            data['surname'],
            data['name']
        )
    
    await message.answer("✅ Ваш вопрос принят! Ожидайте ответа.", reply_markup=start_keyboard)
    await state.finish()

# Обработчики для админов (остаются без изменений)
@dp.message_handler(commands=['answer'], user_id=ADMINS)
async def start_answer_process(message: types.Message):
    questions = get_pending_questions()
    
    if not questions:
        await message.answer("Нет вопросов для ответа", reply_markup=start_keyboard)
        return
    
    keyboard = types.InlineKeyboardMarkup()
    for q in questions:
        keyboard.add(types.InlineKeyboardButton(
            text=f"Вопрос от {q['name']} {q['surname']} ({q['time']})",
            callback_data=f"answer_{q['id']}"
        ))
    
    await message.answer("Выберите вопрос для ответа:", reply_markup=keyboard)
    await AdminActions.select_question.set()

@dp.callback_query_handler(lambda c: c.data.startswith('answer_'), state=AdminActions.select_question)
async def select_question(callback: types.CallbackQuery, state: FSMContext):
    user_id = int(callback.data.split('_')[1])
    questions = get_pending_questions()
    question = next((q for q in questions if q['id'] == user_id), None)
    
    if not question:
        await callback.message.edit_text("Вопрос уже был отвечен")
        return
    
    await state.update_data(
        question_id=user_id,
        question_text=question['question'],
        user_info=question['user_info']
    )
    
    await callback.message.edit_text(
        f"✍️ Вопрос от {question['user_info']}:\n\n"
        f"{question['question']}\n\n"
        f"Введите ваш ответ:"
    )
    await AdminActions.enter_answer.set()

@dp.message_handler(state=AdminActions.enter_answer, user_id=ADMINS)
async def receive_answer(message: types.Message, state: FSMContext):
    data = await state.get_data()
    
    confirm_keyboard = types.InlineKeyboardMarkup()
    confirm_keyboard.row(
        types.InlineKeyboardButton("✅ Подтвердить", callback_data="confirm_answer"),
        types.InlineKeyboardButton("❌ Отменить", callback_data="cancel_answer")
    )
    
    await state.update_data(answer_text=message.text)
    await message.answer(
        f"📝 Ваш ответ:\n\n{message.text}\n\n"
        f"Отправить пользователю {data['user_info']}?",
        reply_markup=confirm_keyboard
    )
    await AdminActions.confirm_answer.set()

@dp.callback_query_handler(lambda c: c.data == "confirm_answer", state=AdminActions.confirm_answer)
async def process_confirm_answer(callback: types.CallbackQuery, state: FSMContext):
    data = await state.get_data()
    
    if not update_question(data['question_id'], data['answer_text'], callback.from_user.id):
        await callback.message.edit_text("Ошибка сохранения")
        await state.finish()
        return
    
    try:
        await bot.send_message(
            chat_id=data['question_id'],
            text=f"📬 Вы получили ответ на ваш вопрос:\n\n"
                 f"{data['question_text']}\n\n"
                 f"💬 Ответ:\n{data['answer_text']}",
            reply_markup=start_keyboard
        )
        await callback.message.edit_text("✅ Ответ отправлен!")
    except Exception as e:
        await callback.message.edit_text(f"Ответ сохранен, но не отправлен: {e}")
    
    await state.finish()

@dp.callback_query_handler(lambda c: c.data == "cancel_answer", state=AdminActions.confirm_answer)
async def process_cancel_answer(callback: types.CallbackQuery, state: FSMContext):
    await callback.message.edit_text("Ответ отменен")
    await state.finish()

# Запуск бота
if __name__ == '__main__':
    from aiogram import executor
    executor.start_polling(dp, skip_updates=True)
